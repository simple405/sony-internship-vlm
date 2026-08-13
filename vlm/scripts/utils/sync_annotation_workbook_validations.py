"""Add standard visible/status dropdowns to annotation template workbooks."""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

from vlm.scripts.utils.atomic_rule_xlsx import (
    ANNOTATION_STATUS_VALUES,
    XLSX_COLUMNS,
    add_annotation_dropdowns,
)


DEFAULT_ROOT = Path("vlm/data/multi_view试标数据集_rev")
REQUIRED_COLUMNS = set(XLSX_COLUMNS[4:10])


def sync_workbook(path: Path) -> None:
    """Refresh annotation validations in one workbook."""
    workbook = openpyxl.load_workbook(path)
    worksheet = workbook.active
    before_rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    headers = {str(value).strip() for value in before_rows[0] if value is not None} if before_rows else set()
    if not REQUIRED_COLUMNS <= headers:
        raise ValueError(f"{path}: not a visible/status annotation workbook")

    add_annotation_dropdowns(worksheet)
    temporary_path = path.with_suffix(".validation.tmp.xlsx")
    workbook.save(temporary_path)
    workbook.close()

    verification = openpyxl.load_workbook(temporary_path, read_only=False, data_only=False)
    checked_sheet = verification.active
    after_rows = [tuple(row) for row in checked_sheet.iter_rows(values_only=True)]
    formulas = {validation.formula1 for validation in checked_sheet.data_validations.dataValidation}
    verification.close()
    if after_rows != before_rows:
        raise ValueError(f"{path}: validation sync changed worksheet values")
    if not any("\u4f4d\u7f6e\u9519\u8bef" in formula for formula in formulas):
        raise ValueError(f"{path}: missing \u4f4d\u7f6e\u9519\u8bef status validation")
    temporary_path.replace(path)


def main() -> None:
    """Refresh validations across generated annotation workbooks."""
    parser = argparse.ArgumentParser(description="Sync annotation workbook dropdowns in place.")
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    args = parser.parse_args()
    paths = sorted(
        path
        for path in args.root.rglob("*.xlsx")
        if ".tmp." not in path.name and not path.name.startswith("~$")
    )
    if not paths:
        raise SystemExit(f"No .xlsx files found under {args.root}")
    for path in paths:
        sync_workbook(path)
    print(f"synced_workbooks={len(paths)} status_values={','.join(ANNOTATION_STATUS_VALUES)}")


if __name__ == "__main__":
    main()

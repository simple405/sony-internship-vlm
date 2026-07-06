"""Export a formatted Excel workbook for manual supervision review."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


DEFAULT_REPORT_ROOT = Path(
    "vlm/data/safebooru_2d/japanese_anime_turnaround_pilot_20/reports/supervision_review/runs"
)

REVIEW_COLUMNS = [
    "human_decision",
    "human_notes",
    "correct_findings",
    "missed_findings",
    "false_positive_findings",
    "sample_id",
    "category",
    "decision",
    "expected_decision",
    "score",
    "primary_issue",
    "agent_summary",
    "source_image_path",
    "generated_image_path",
]

DECISION_VALUES = ["approved", "needs_revision", "rejected", "unclear"]
FINDING_TYPES = [
    "identity_missing",
    "identity_hallucination",
    "misread_accessory",
    "hairstyle_structure_error",
    "face_expression_error",
    "product_type_error",
    "view_inconsistency",
    "pattern_topology_error",
    "layout_error",
    "safety_or_generation_block",
]


def parse_args() -> argparse.Namespace:
    # Note 1: This exporter fixes the common CSV usability problem: raw CSV has
    # no frozen header, widths, wrapping, colors, or dropdowns, so long path
    # fields make columns look misaligned in text editors.
    parser = argparse.ArgumentParser(description="Export a formatted manual-review workbook from human_review_queue.csv.")
    parser.add_argument("--human-review-csv", type=Path, help="Path to human_review_queue.csv.")
    parser.add_argument("--run-dir", type=Path, help="Run directory containing tables/human_review_queue.csv.")
    parser.add_argument("--output", type=Path, help="Output .xlsx path.")
    return parser.parse_args()


def resolve_input(args: argparse.Namespace) -> Path:
    if args.human_review_csv:
        return args.human_review_csv
    if args.run_dir:
        return args.run_dir / "tables" / "human_review_queue.csv"
    candidates = sorted(
        DEFAULT_REPORT_ROOT.glob("*/tables/human_review_queue.csv"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise SystemExit("No human_review_queue.csv found. Pass --human-review-csv or --run-dir.")
    return candidates[0]


def resolve_output(args: argparse.Namespace, input_path: Path) -> Path:
    if args.output:
        return args.output
    return input_path.parent / "human_review_queue.xlsx"


def read_csv(path: Path) -> list[dict[str, str]]:
    # Note 2: utf-8-sig reads both normal UTF-8 and Excel-saved UTF-8 with BOM.
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def ordered_row(row: dict[str, str]) -> list[str]:
    return [row.get(column, "") for column in REVIEW_COLUMNS]


def style_review_sheet(ws, row_count: int) -> None:
    # Note 3: Freeze F2 keeps row 1 and the fillable columns A:E visible while
    # scrolling through long source/generated paths.
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = True

    header_fill = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    readonly_fill = PatternFill("solid", fgColor="EAF2F8")
    header_font = Font(bold=True, color="FFFFFF")

    widths = {
        "A": 18,
        "B": 48,
        "C": 28,
        "D": 28,
        "E": 30,
        "F": 14,
        "G": 22,
        "H": 18,
        "I": 20,
        "J": 10,
        "K": 34,
        "L": 56,
        "M": 72,
        "N": 72,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=max(2, row_count + 1), min_col=1, max_col=len(REVIEW_COLUMNS)):
        for index, cell in enumerate(row, start=1):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = input_fill if index <= 5 else readonly_fill
        ws.row_dimensions[cell.row].height = 58

    for header_cell in ws[1]:
        if header_cell.value == "human_decision":
            header_cell.comment = Comment("Use: approved, needs_revision, rejected, or unclear.", "Codex")
        elif header_cell.value in {"correct_findings", "missed_findings", "false_positive_findings"}:
            header_cell.comment = Comment("Use semicolon-separated finding types, for example identity_missing;layout_error.", "Codex")

    if row_count > 0:
        table_ref = f"A1:{get_column_letter(len(REVIEW_COLUMNS))}{row_count + 1}"
        table = Table(displayName="HumanReviewQueue", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)

        validation = DataValidation(type="list", formula1='"approved,needs_revision,rejected,unclear"', allow_blank=True)
        validation.error = "Choose one of: approved, needs_revision, rejected, unclear."
        validation.errorTitle = "Invalid human_decision"
        ws.add_data_validation(validation)
        validation.add(f"A2:A{row_count + 1}")


def add_instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions")
    ws["A1"] = "Manual Review Workbook"
    ws["A1"].font = Font(bold=True, size=14)
    instructions = [
        ("human_decision", "approved / needs_revision / rejected / unclear"),
        ("human_notes", "Short Chinese note explaining the decision."),
        ("correct_findings", "Finding types the agent correctly reported. Empty is OK for dry-run."),
        ("missed_findings", "Finding types you found but the agent missed."),
        ("false_positive_findings", "Finding types the agent reported but you think are wrong."),
        ("separator", "Use semicolon to separate multiple finding types."),
    ]
    for row_index, (field, meaning) in enumerate(instructions, start=3):
        ws.cell(row=row_index, column=1, value=field)
        ws.cell(row=row_index, column=2, value=meaning)
    ws["A11"] = "Allowed finding types"
    ws["A11"].font = Font(bold=True)
    for row_index, finding_type in enumerate(FINDING_TYPES, start=12):
        ws.cell(row=row_index, column=1, value=finding_type)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 72
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_workbook(input_path: Path, output_path: Path) -> dict[str, str | int]:
    rows = read_csv(input_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Human Review"
    ws.append(REVIEW_COLUMNS)
    for row in rows:
        ws.append(ordered_row(row))
    style_review_sheet(ws, len(rows))
    add_instructions_sheet(wb)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {"input": str(input_path), "output": str(output_path), "rows": len(rows)}


def main() -> None:
    args = parse_args()
    input_path = resolve_input(args)
    if not input_path.exists():
        raise SystemExit(f"human review CSV does not exist: {input_path}")
    output_path = resolve_output(args, input_path)
    summary = export_workbook(input_path, output_path)
    print(json.dumps({"status": "finished", **summary}, ensure_ascii=False))


if __name__ == "__main__":
    main()

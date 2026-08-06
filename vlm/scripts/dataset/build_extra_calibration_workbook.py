"""Create an Excel review sheet from visual-verifier JSONL results."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_INPUT = Path("vlm/data/1-动漫标注结果导出_paired_samples/extra_prediction_visual_verification_holdout30_v1.jsonl")
DEFAULT_OUTPUT = Path("vlm/data/1-动漫标注结果导出_paired_samples/extra_prediction_calibration_holdout30_v1.xlsx")
DEFAULT_PAIRED_ROOT = Path("vlm/data/1-动漫标注结果导出_paired_samples")
DEFAULT_CROP_ROOT = Path("vlm/tmp/extra_prediction_verifier_crops")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]


def build_workbook(rows: list[dict[str, Any]], *, paired_root: Path, crop_root: Path) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "extra_calibration"
    headers = [
        "sample_id", "strict_extra_index", "source_image", "crop_image", "crop_preview",
        "prediction_element", "prediction_description", "bbox",
        "verifier_image_grounded", "verifier_description_correct", "verifier_duplicate", "verifier_conflict",
        "human_image_grounded", "human_description_correct", "human_duplicate", "human_conflict",
        "human_accepted_extra", "notes",
    ]
    sheet.append(headers)
    for row in rows:
        prediction = row.get("prediction", {})
        verifier = row.get("visual_verifier", {})
        sample_id = str(row.get("sample_id", ""))
        extra_index = int(row.get("strict_extra_index", -1))
        crop_path = f"vlm/tmp/extra_prediction_verifier_crops/{sample_id}/extra_{extra_index:03d}.png"
        crop_file = crop_root / sample_id / f"extra_{extra_index:03d}.png"
        source_file = paired_root / str(row.get("source_image", ""))
        sheet.append(
            [
                sample_id,
                extra_index,
                row.get("source_image", ""),
                crop_path,
                None,
                prediction.get("element", ""),
                prediction.get("description", ""),
                json.dumps(prediction.get("bbox", []), ensure_ascii=False),
                verifier.get("image_grounded"),
                verifier.get("description_correct"),
                verifier.get("duplicate"),
                verifier.get("conflict"),
                None, None, None, None,
                None,
                row.get("notes", ""),
            ]
        )
        row_number = sheet.max_row
        for column, target in ((3, source_file), (4, crop_file)):
            if target.is_file():
                sheet.cell(row_number, column).hyperlink = target.resolve().as_uri()
                sheet.cell(row_number, column).style = "Hyperlink"
        if crop_file.is_file():
            preview = ExcelImage(crop_file)
            scale = min(160 / preview.width, 100 / preview.height, 1.0)
            preview.width *= scale
            preview.height *= scale
            sheet.add_image(preview, f"E{row_number}")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 20, "B": 12, "C": 46, "D": 56, "E": 24, "F": 22, "G": 58, "H": 28, "I": 18, "J": 24, "K": 18, "L": 18, "M": 18, "N": 24, "O": 18, "P": 18, "Q": 20, "R": 30}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row_cells[0].row].height = 82

    boolean_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    boolean_validation.error = "请选择 TRUE 或 FALSE"
    boolean_validation.errorTitle = "布尔值无效"
    sheet.add_data_validation(boolean_validation)
    boolean_validation.add(f"M2:P{len(rows) + 1}")
    for row_number in range(2, len(rows) + 2):
        # Human accepted_extra is derived after all four human booleans exist.
        sheet.cell(row_number, 17).value = (
            f'=IF(COUNTBLANK(M{row_number}:P{row_number})>0,"",'
            f'AND(M{row_number}=TRUE,N{row_number}=TRUE,O{row_number}=FALSE,P{row_number}=FALSE))'
        )
    sheet.conditional_formatting.add(
        f"Q2:Q{len(rows) + 1}",
        CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor="C6EFCE")),
    )
    return workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build an Excel sheet for manual extra calibration.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--paired-root", type=Path, default=DEFAULT_PAIRED_ROOT)
    parser.add_argument("--crop-root", type=Path, default=DEFAULT_CROP_ROOT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rows = read_jsonl(args.input)
    if not rows:
        raise SystemExit("input JSONL is empty")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(rows, paired_root=args.paired_root.resolve(), crop_root=args.crop_root.resolve()).save(args.output)
    print(json.dumps({"output": str(args.output), "row_count": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()

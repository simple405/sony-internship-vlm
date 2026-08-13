"""Audit translation caches and XLSX value cells for forbidden dunhao punctuation."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from vlm.scripts.utils.atomic_rule_xlsx import TRANSLATION_CACHE_RELATIVE_PATH


DEFAULT_DATASET = Path("vlm/data/design_sheet_10610_smoke30")


def parse_args() -> argparse.Namespace:
    """Parse Chinese-value audit arguments."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset-dir", type=Path, default=DEFAULT_DATASET)
    return parser.parse_args()


def audit_cache(path: Path) -> list[dict[str, Any]]:
    """Return cache entries whose ``value_cn`` contains a dunhao."""
    if not path.is_file():
        return []
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    translations = payload.get("translations", {})
    if not isinstance(translations, dict):
        raise ValueError(f"translations must be an object: {path}")
    findings: list[dict[str, Any]] = []
    for key, entry in translations.items():
        if not isinstance(entry, dict):
            continue
        value_cn = str(entry.get("value_cn", ""))
        if "、" in value_cn:
            findings.append(
                {
                    "source": "cache",
                    "path": str(path),
                    "key": str(key),
                    "rule_id": str(entry.get("rule_id", "")),
                    "value": str(entry.get("value", "")),
                    "value_cn": value_cn,
                }
            )
    return findings


def audit_workbook(path: Path) -> list[dict[str, Any]]:
    """Return XLSX cells in the Chinese value column that contain a dunhao."""
    findings: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(value or "") for value in first_row]
        value_column = headers.index("取值") + 1 if "取值" in headers else 4
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                min_col=value_column,
                max_col=value_column,
                values_only=True,
            ),
            start=2,
        ):
            value_cn = "" if row[0] is None else str(row[0])
            if "、" in value_cn:
                findings.append(
                    {
                        "source": "xlsx",
                        "path": str(path),
                        "sheet": sheet.title,
                        "cell": f"{sheet.cell(row_number, value_column).coordinate}",
                        "value_cn": value_cn,
                    }
                )
    finally:
        workbook.close()
    return findings


def audit_dataset(dataset_dir: Path) -> dict[str, Any]:
    """Audit the dataset cache and every XLSX file below the dataset root."""
    cache_path = dataset_dir / TRANSLATION_CACHE_RELATIVE_PATH
    workbooks = sorted(dataset_dir.rglob("*.xlsx"))
    findings = audit_cache(cache_path)
    workbook_errors: list[dict[str, str]] = []
    for workbook in workbooks:
        try:
            findings.extend(audit_workbook(workbook))
        except Exception as exc:  # noqa: BLE001
            workbook_errors.append(
                {
                    "path": str(workbook),
                    "error_type": type(exc).__name__,
                    "error": str(exc),
                }
            )
    return {
        "status": "passed" if not findings and not workbook_errors else "failed",
        "dataset_dir": str(dataset_dir),
        "cache_path": str(cache_path),
        "xlsx_files_scanned": len(workbooks),
        "findings": findings,
        "workbook_errors": workbook_errors,
    }


def main() -> None:
    """Print a machine-readable audit and fail when forbidden values are found."""
    args = parse_args()
    report = audit_dataset(args.dataset_dir.resolve())
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
    if report["status"] != "passed":
        raise SystemExit(1)


if __name__ == "__main__":
    main()

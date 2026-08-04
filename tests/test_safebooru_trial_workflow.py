"""Focused tests for the Safebooru trial dataset workflow."""

import json
from pathlib import Path

from openpyxl import load_workbook

from vlm.scripts.build_safebooru_trial_dataset import (
    HEAD_ONLY_CATEGORIES,
    find_multiview_file,
    write_xlsx,
)
from vlm.scripts.data.assign_merchandise_categories import (
    CATEGORIES,
    apply_balanced_assignment,
    balanced_target_counts,
)
from vlm.scripts.extract_atomic_rules_safebooru import normalize_rules, parse_json
from vlm.scripts.utils.tag_and_export_atomic_rules import write_xlsx as export_write_xlsx
from vlm.scripts.utils.sync_generated_annotation_xlsx import sync_sample


def test_parse_json_accepts_fenced_model_output():
    parsed = parse_json('```json\n{"code":"1","atomic_rules":[]}\n```')
    assert parsed["code"] == "1"


def test_normalize_rules_deduplicates_and_keeps_scalar_values(tmp_path: Path):
    image = tmp_path / "1.jpg"
    image.write_bytes(b"image")
    result = normalize_rules(
        {"atomic_rules": [
            {"id": "Hair Color", "location": "head", "value": "orange"},
            {"id": "hair_color", "location": "head", "value": "duplicate"},
            {"id": "has_glasses", "location": "head", "value": True},
        ]},
        "1",
        image,
        "qwen3-vl-plus",
        "prompt",
        0.1,
    )
    assert [rule["id"] for rule in result["atomic_rules"]] == ["hair_color", "has_glasses"]
    assert result["atomic_rules"][0]["location"] == "head"
    assert result["atomic_rules"][1]["value"] is True


def test_head_only_xlsx_filters_body_rules(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "hair_color", "value": "orange"},
            {"id": "top_color", "value": "blue"},
        ],
    }), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_xlsx(json_path, output_path, {"hair_color": "head", "top_color": "body"}, HEAD_ONLY_CATEGORIES[0])
    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "sample_id", "rule_id", "location", "value", "front_visible",
        "front_status", "side_visible", "side_status", "back_visible",
        "back_status", "note",
    )
    assert len(rows) == 2
    assert rows[1][1:4] == ("hair_color", "head", "orange")


def test_xlsx_status_dropdown_includes_wrong_prosition(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({"code": "123", "atomic_rules": []}), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_xlsx(json_path, output_path, {}, "plush")
    workbook = load_workbook(output_path, read_only=False)
    formulas = {validation.formula1 for validation in workbook.active.data_validations.dataValidation}
    workbook.close()

    assert any("wrong_prosition" in formula for formula in formulas)


def test_sync_generated_sample_writes_xlsx_after_image_download(tmp_path: Path):
    sample_dir = tmp_path / "plush" / "123"
    sample_dir.mkdir(parents=True)
    (sample_dir / "123_atomic_rules.json").write_text(json.dumps({
        "code": "123",
        "atomic_rules": [{"id": "hair_color", "location": "head", "value": "red"}],
    }), encoding="utf-8")
    (sample_dir / "123_plush.png").write_bytes(b"image")

    assert sync_sample(sample_dir, "plush", "plush") is True
    assert (sample_dir / "123.xlsx").exists()
    assert sync_sample(sample_dir, "plush", "plush") is False


def test_position_rules_keep_emitted_annotator_view(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "horn_position", "value": "right"},
            {"id": "armband_position", "value": "left_arm"},
            {"id": "wing_position", "value": "both_sides"},
        ],
    }), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_xlsx(json_path, output_path, {"horn_position": "body", "wing_position": "body"}, "plush")
    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][1:4] == ("horn_position", "body", "right")
    assert rows[2][1:4] == ("armband_position", "body", "left_arm")
    assert rows[3][1:4] == ("wing_position", "body", "both_sides")


def test_generated_head_only_paths_filter_body_rules(tmp_path: Path):
    sample_dir = tmp_path / "head_key_chain" / "123"
    json_path = sample_dir / "123_atomic_rules.json"
    sample_dir.mkdir(parents=True)
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "hair_color", "value": "orange"},
            {"id": "top_color", "value": "blue"},
        ],
    }), encoding="utf-8")
    output_path = export_write_xlsx(json_path, {"hair_color": "head", "top_color": "body"})
    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][1:4] == ("hair_color", "head", "orange")


def test_balanced_assignment_preserves_existing_and_fills_targets():
    rows = []
    targets = balanced_target_counts(12)
    for index in range(12):
        existing = "head_key_chain" if index < targets["head_key_chain"] else ""
        row = {
            "sample_id": f"{index:03d}",
            "already_generated_category": existing,
            "primary_category": existing or "head_key_chain",
            "assignment_source": "existing_generated_output" if existing else "rule_score",
            "candidate_categories": "|".join(CATEGORIES),
            "primary_score": 100,
        }
        for category in CATEGORIES:
            row[f"score_{category}"] = 100 if category == "head_key_chain" else 10
        rows.append(row)
    apply_balanced_assignment(rows)
    counts = {category: sum(row["primary_category"] == category for row in rows) for category in CATEGORIES}
    assert counts == targets
    assert all(row["primary_category"] == "head_key_chain" for row in rows[:targets["head_key_chain"]])


def test_find_multiview_file_prefers_canonical_category_suffix(tmp_path: Path):
    sample_dir = tmp_path / "dataset_figurine" / "123"
    sample_dir.mkdir(parents=True)
    fallback = sample_dir / "123_figurine_current_prompt_rerun.png"
    canonical = sample_dir / "123_figurine.png"
    fallback.write_bytes(b"fallback")
    canonical.write_bytes(b"canonical")
    assert find_multiview_file(tmp_path, "dataset_figurine", "123") == canonical

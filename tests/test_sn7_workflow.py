"""Focused tests for the retained SN-7 dataset workflow."""

import json
from argparse import Namespace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from vlm.scripts.import_sn7_design_sheet_dataset import import_image
from vlm.scripts.package_sn7_dataset import find_atomic_file, find_multiview_file, package_dataset
from vlm.scripts.generate import generate_sn7_multiview
from vlm.scripts.data.assign_merchandise_categories import (
    CATEGORIES,
    apply_balanced_assignment,
    balanced_target_counts,
    discover_generated,
)
from vlm.scripts.extract_atomic_rules import is_probable_black_white_line_art, normalize_rules, parse_json
from vlm.scripts.utils.atomic_rule_xlsx import display_rule_id, display_value, write_atomic_rules_xlsx
from vlm.scripts.utils.sync_annotation_workbook_validations import sync_workbook
from vlm.scripts.utils.sync_generated_annotation_xlsx import has_generated_image, sync_sample


def test_parse_json_accepts_fenced_model_output():
    parsed = parse_json('```json\n{"code":"1","atomic_rules":[]}\n```')
    assert parsed["code"] == "1"


def test_package_atomic_error_wins_over_stale_rules(tmp_path: Path):
    sample_root = tmp_path / "atomic_rules" / "sample"
    sample_root.mkdir(parents=True)
    (sample_root / "atomic_rules.json").write_text("{}", encoding="utf-8")
    (sample_root / "error.json").write_text("{}", encoding="utf-8")

    assert find_atomic_file(tmp_path / "atomic_rules", "sample") is None


def test_normalize_rules_deduplicates_and_keeps_scalar_values(tmp_path: Path):
    image = tmp_path / "1.jpg"
    image.write_bytes(b"image")
    result = normalize_rules(
        {"atomic_rules": [
            {"id": "Hair Color", "location": "head", "value": "orange"},
            {"id": "hair_color", "location": "head", "value": "duplicate"},
            {"id": "has_glasses", "location": "head", "value": True},
        ]},
        "1",
        image,
        "qwen3-vl-plus",
        "prompt",
        0.1,
    )
    assert [rule["id"] for rule in result["atomic_rules"]] == ["hair_color", "has_glasses"]
    assert result["atomic_rules"][0]["location"] == "头部"
    assert result["atomic_rules"][1]["value"] is True


def test_normalize_rules_preserves_chinese_contract(tmp_path: Path):
    image = tmp_path / "1.jpg"
    image.write_bytes(b"image")
    result = normalize_rules(
        {"atomic_rules": [
            {"id": "发色", "location": "头部", "value": "金色"},
            {"id": "上衣类型", "location": "身体", "value": "黑色外套"},
        ]},
        "1",
        image,
        "qwen3-vl-plus",
        "prompt",
        0.1,
    )

    assert result["atomic_rules"] == [
        {"id": "发色", "location": "头部", "value": "金色"},
        {"id": "上衣类型", "location": "身体", "value": "黑色外套"},
    ]


def test_normalize_rules_rejects_model_line_art_error(tmp_path: Path):
    image = tmp_path / "1.jpg"
    image.write_bytes(b"image")

    with pytest.raises(ValueError, match="black_white_line_art"):
        normalize_rules(
            {
                "error": {
                    "type": "black_white_line_art",
                    "message": "输入图是明显的黑白线条稿，跳过 atomic_rules 提取",
                }
            },
            "1",
            image,
            "qwen3-vl-plus",
            "prompt",
            0.1,
        )


def test_probable_black_white_line_art_detector_is_conservative(tmp_path: Path):
    line_art = tmp_path / "line_art.png"
    image = Image.new("RGB", (120, 120), "white")
    pixels = image.load()
    for offset in range(20, 100):
        pixels[offset, 20] = (0, 0, 0)
        pixels[offset, 99] = (0, 0, 0)
        pixels[20, offset] = (0, 0, 0)
        pixels[99, offset] = (0, 0, 0)
    image.save(line_art)

    colored = tmp_path / "colored.png"
    Image.new("RGB", (120, 120), "orange").save(colored)

    assert is_probable_black_white_line_art(line_art) is True
    assert is_probable_black_white_line_art(colored) is False


def test_head_only_xlsx_filters_body_rules(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "hair_color", "value": "orange"},
            {"id": "top_color", "value": "blue"},
        ],
    }), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_atomic_rules_xlsx(
        json_path,
        output_path,
        "head_key_chain",
        {"hair_color": "head", "top_color": "body"},
    )
    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "样本编号", "规则名称", "位置", "取值", "正面可见性",
        "正面状态", "侧面可见性", "侧面状态", "背面可见性",
        "\u80cc\u9762\u72b6\u6001", "\u5907\u6ce8", "\u539f\u59cb\u89c4\u5219ID", "\u539f\u59cb\u53d6\u503c",
    )
    assert len(rows) == 2
    assert rows[1][1:4] == ("头发、颜色", "头部", "橙色")


def test_head_only_xlsx_filters_chinese_body_rules(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "发色", "location": "头部", "value": "金色"},
            {"id": "上衣类型", "location": "身体", "value": "黑色外套"},
        ],
    }, ensure_ascii=False), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"

    write_atomic_rules_xlsx(json_path, output_path, "head_key_chain")

    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][1:4] == ("发色", "头部", "金色")


def test_xlsx_status_dropdown_includes_wrong_prosition(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({"code": "123", "atomic_rules": []}), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_atomic_rules_xlsx(json_path, output_path, "plush")
    workbook = load_workbook(output_path, read_only=False)
    formulas = {validation.formula1 for validation in workbook.active.data_validations.dataValidation}
    workbook.close()

    assert any("位置错误" in formula for formula in formulas)



def test_sync_annotation_workbook_validations_accepts_current_chinese_schema(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({"code": "123", "atomic_rules": []}), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_atomic_rules_xlsx(json_path, output_path, "plush")

    sync_workbook(output_path)

    workbook = load_workbook(output_path, read_only=False)
    formulas = {validation.formula1 for validation in workbook.active.data_validations.dataValidation}
    workbook.close()
    assert any("\u4f4d\u7f6e\u9519\u8bef" in formula for formula in formulas)


def test_sync_generated_sample_writes_xlsx_after_image_download(tmp_path: Path):
    sample_dir = tmp_path / "plush" / "123"
    sample_dir.mkdir(parents=True)
    (sample_dir / "123_atomic_rules.json").write_text(json.dumps({
        "code": "123",
        "atomic_rules": [{"id": "hair_color", "location": "head", "value": "red"}],
    }), encoding="utf-8")
    Image.new("RGB", (8, 8), "white").save(sample_dir / "123_plush.png")

    assert sync_sample(sample_dir, "plush", "plush") is True
    assert (sample_dir / "123.xlsx").exists()
    assert sync_sample(sample_dir, "plush", "plush") is False
    assert sync_sample(sample_dir, "plush", "plush", overwrite=True) is True


def test_position_rules_keep_emitted_annotator_view(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "horn_position", "location": "body", "value": "right"},
            {"id": "armband_position", "location": "body", "value": "left_arm"},
            {"id": "wing_position", "location": "body", "value": "both_sides"},
        ],
    }), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    write_atomic_rules_xlsx(
        json_path,
        output_path,
        "plush",
        {},
    )
    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][1:4] == ("角、位置", "身体", "右")
    assert rows[2][1:4] == ("臂带、位置", "身体", "左手臂")
    assert rows[3][1:4] == ("翅膀、位置", "身体", "双侧侧面")



def test_unknown_xlsx_tokens_keep_raw_evidence():
    assert display_rule_id("gemstone_color") == "\u672a\u7ffb\u8bd1:gemstone\u3001\u989c\u8272"
    assert display_value("cerulean_teal") == "\u672a\u7ffb\u8bd1:cerulean\u84dd\u7eff\u8272"


def test_xlsx_uses_contextual_qwen_translation_cache(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "skirt_length", "location": "body", "value": "knee_length"},
        ],
    }), encoding="utf-8")
    output_path = tmp_path / "123.xlsx"
    from vlm.scripts.utils.atomic_rule_xlsx import translation_key

    key = translation_key("skirt_length", "knee_length")
    translations = {
        key: {
            "rule_id": "skirt_length",
            "value": "knee_length",
            "rule_name_cn": "裙长",
            "value_cn": "及膝长度",
        }
    }

    write_atomic_rules_xlsx(
        json_path,
        output_path,
        "plush",
        translation_cache=translations,
    )

    sheet = load_workbook(output_path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][1:4] == ("裙长", "身体", "及膝长度")
    assert rows[1][11:13] == ("skirt_length", "knee_length")


def test_xlsx_rejects_dunhao_in_contextual_chinese_value(tmp_path: Path):
    json_path = tmp_path / "atomic_rules.json"
    json_path.write_text(json.dumps({
        "code": "123",
        "atomic_rules": [
            {"id": "bag_color", "location": "body", "value": "white_and_green"},
        ],
    }), encoding="utf-8")
    from vlm.scripts.utils.atomic_rule_xlsx import translation_key

    translations = {
        translation_key("bag_color", "white_and_green"): {
            "rule_id": "bag_color",
            "value": "white_and_green",
            "rule_name_cn": "包袋配色",
            "value_cn": "白色、绿色",
        }
    }

    with pytest.raises(ValueError, match="forbidden dunhao"):
        write_atomic_rules_xlsx(
            json_path,
            tmp_path / "123.xlsx",
            "plush",
            translation_cache=translations,
        )


def test_zero_byte_generated_image_is_not_complete(tmp_path: Path):
    sample_dir = tmp_path / "plush" / "sample"
    sample_dir.mkdir(parents=True)
    (sample_dir / "sample_plush.png").write_bytes(b"")

    assert has_generated_image(sample_dir, "sample", "plush") is False
    assert find_multiview_file(tmp_path, "plush", "sample") is None


def test_sn7_output_suffix_defaults_to_category_mapping(monkeypatch, tmp_path: Path):
    prompt = tmp_path / "prompt.txt"
    prompt.write_text("{{MERCHANDISE_CATEGORY}} {{VIEW_REQUIREMENTS}} {{CATEGORY_REQUIREMENTS}}", encoding="utf-8")
    image = tmp_path / "source" / "sample.png"
    image.parent.mkdir()
    Image.new("RGB", (8, 8), "white").save(image)
    atomic = tmp_path / "atomic" / "sample" / "atomic_rules.json"
    atomic.parent.mkdir(parents=True)
    atomic.write_text(json.dumps({"code": "sample", "atomic_rules": []}), encoding="utf-8")
    args = Namespace(
        sample_id=["sample"],
        source_dir=image.parent,
        atomic_dir=tmp_path / "atomic",
        direct_output_dir=tmp_path / "out",
        prompt_file=prompt,
        aspect_ratio="21:9",
        resolution="1k",
        output_suffix=None,
        category="plush",
        poll_interval=1,
        timeout=1,
        workers=1,
        stop_on_error=False,
        keep_debug_files=False,
        dry_run=True,
    )
    monkeypatch.setattr(generate_sn7_multiview, "parse_args", lambda: args)

    generate_sn7_multiview.main()

    assert args.output_suffix == "plush"


def test_sn7_output_suffix_rejects_category_mismatch(monkeypatch, tmp_path: Path):
    prompt = tmp_path / "prompt.txt"
    prompt.write_text("{{MERCHANDISE_CATEGORY}} {{VIEW_REQUIREMENTS}} {{CATEGORY_REQUIREMENTS}}", encoding="utf-8")
    args = Namespace(
        sample_id=["sample"],
        source_dir=tmp_path,
        atomic_dir=tmp_path,
        direct_output_dir=tmp_path / "out",
        prompt_file=prompt,
        aspect_ratio="21:9",
        resolution="1k",
        output_suffix="head_keychain",
        category="plush",
        poll_interval=1,
        timeout=1,
        workers=1,
        stop_on_error=False,
        keep_debug_files=False,
        dry_run=True,
    )
    monkeypatch.setattr(generate_sn7_multiview, "parse_args", lambda: args)

    with pytest.raises(ValueError, match="does not match category"):
        generate_sn7_multiview.main()


def test_runninghub_run_one_skips_existing_generated_output_without_overwrite(tmp_path: Path):
    source = tmp_path / "source" / "sample.png"
    source.parent.mkdir()
    Image.new("RGB", (8, 8), "orange").save(source)
    atomic = tmp_path / "atomic" / "sample" / "atomic_rules.json"
    atomic.parent.mkdir(parents=True)
    atomic.write_text(json.dumps({"code": "sample", "atomic_rules": []}), encoding="utf-8")
    clean_dir = tmp_path / "out" / "sample"
    clean_dir.mkdir(parents=True)
    Image.new("RGB", (8, 8), "white").save(clean_dir / "sample_plush.png")
    sentinel = clean_dir / "sample_atomic_rules.json"
    sentinel.write_text("previous", encoding="utf-8")
    args = Namespace(
        source_dir=source.parent,
        atomic_dir=tmp_path / "atomic",
        direct_output_dir=tmp_path / "out",
        prompt_file=tmp_path / "prompt.txt",
        aspect_ratio="21:9",
        resolution="1k",
        output_suffix="plush",
        category="plush",
        poll_interval=1,
        timeout=1,
        workers=1,
        stop_on_error=False,
        keep_debug_files=False,
        skip_xlsx=True,
        dry_run=False,
    )

    status = generate_sn7_multiview.run_one(args, "prompt", "api-key", "sample")

    assert status["status"] == "skipped_existing"
    assert status["existing_output"].endswith("sample_plush.png")
    assert sentinel.read_text(encoding="utf-8") == "previous"


def test_runninghub_run_one_rejects_black_white_line_art_before_submission(tmp_path: Path):
    source = tmp_path / "source" / "sample.png"
    source.parent.mkdir()
    image = Image.new("RGB", (120, 120), "white")
    pixels = image.load()
    for offset in range(20, 100):
        pixels[offset, 20] = (0, 0, 0)
        pixels[offset, 99] = (0, 0, 0)
        pixels[20, offset] = (0, 0, 0)
        pixels[99, offset] = (0, 0, 0)
    image.save(source)
    atomic = tmp_path / "atomic" / "sample" / "atomic_rules.json"
    atomic.parent.mkdir(parents=True)
    atomic.write_text(json.dumps({"code": "sample", "atomic_rules": []}), encoding="utf-8")
    args = Namespace(
        source_dir=source.parent,
        atomic_dir=tmp_path / "atomic",
        direct_output_dir=tmp_path / "out",
        prompt_file=tmp_path / "prompt.txt",
        aspect_ratio="21:9",
        resolution="1k",
        output_suffix="plush",
        category="plush",
        poll_interval=1,
        timeout=1,
        workers=1,
        stop_on_error=False,
        keep_debug_files=False,
        skip_xlsx=True,
        dry_run=False,
    )

    with pytest.raises(ValueError, match="black_white_line_art"):
        generate_sn7_multiview.run_one(args, "prompt", "api-key", "sample")

    assert not (tmp_path / "out" / "sample").exists()


def test_balanced_assignment_preserves_existing_and_fills_targets():
    rows = []
    targets = balanced_target_counts(12)
    for index in range(12):
        existing = "head_key_chain" if index < targets["head_key_chain"] else ""
        row = {
            "sample_id": f"{index:03d}",
            "already_generated_category": existing,
            "primary_category": existing or "head_key_chain",
            "assignment_source": "existing_generated_output" if existing else "rule_score",
            "candidate_categories": "|".join(CATEGORIES),
            "primary_score": 100,
        }
        for category in CATEGORIES:
            row[f"score_{category}"] = 100 if category == "head_key_chain" else 10
        rows.append(row)
    apply_balanced_assignment(rows)
    counts = {category: sum(row["primary_category"] == category for row in rows) for category in CATEGORIES}
    assert counts == targets
    assert all(row["primary_category"] == "head_key_chain" for row in rows[:targets["head_key_chain"]])


def test_find_multiview_file_prefers_canonical_category_suffix(tmp_path: Path):
    sample_dir = tmp_path / "dataset_figurine" / "123"
    sample_dir.mkdir(parents=True)
    fallback = sample_dir / "123_figurine_current_prompt_rerun.png"
    canonical = sample_dir / "123_figurine.png"
    fallback.write_bytes(b"fallback")
    Image.new("RGB", (8, 8), "white").save(canonical)
    assert find_multiview_file(tmp_path, "dataset_figurine", "123") == canonical


def test_import_image_does_not_replace_valid_target_with_corrupt_input(
    tmp_path: Path,
):
    source = tmp_path / "source.png"
    source.write_bytes(b"not-an-image")
    target = tmp_path / "target.png"
    Image.new("RGB", (8, 8), "white").save(target)
    original = target.read_bytes()

    with pytest.raises(Exception):
        import_image(source, target)

    assert target.read_bytes() == original
    assert not target.with_suffix(".png.part").exists()


def test_incomplete_package_preserves_previous_deliverable(tmp_path: Path):
    root = tmp_path / "dataset"
    image = root / "image" / "sample.png"
    image.parent.mkdir(parents=True)
    Image.new("RGB", (8, 8), "white").save(image)
    (root / "manifest.csv").write_text(
        "post_id,image_path\nsample,image/sample.png\n", encoding="utf-8"
    )
    assignment = root / "assignments.csv"
    assignment.write_text(
        "sample_id,primary_category\nsample,plush\n", encoding="utf-8"
    )
    package_root = root / "deliverables"
    package_root.mkdir()
    sentinel = package_root / "previous.txt"
    sentinel.write_text("previous", encoding="utf-8")
    args = Namespace(
        root=root,
        manifest=None,
        assignment_csv=assignment,
        generated_root=None,
        package_root=package_root,
        overwrite=True,
        allow_incomplete=False,
    )

    summary = package_dataset(args)

    assert summary["incomplete_count"] == 1
    assert summary["package_replaced"] is False
    assert sentinel.read_text(encoding="utf-8") == "previous"
    assert (root / "reports" / "package_validation" / "incomplete_samples.csv").is_file()


def test_category_assignment_ignores_incomplete_generated_directory(tmp_path: Path):
    sample_dir = tmp_path / "generated" / "plush" / "sample"
    sample_dir.mkdir(parents=True)
    (sample_dir / "sample_atomic_rules.json").write_text("{}", encoding="utf-8")

    assert discover_generated(tmp_path) == {}
